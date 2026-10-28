"""Read prediction checkmarks and target markings from the sensor workbook."""
from openpyxl import load_workbook

# Airflow targets follow the headings, explicitly confirmed by the project owner.
TARGETS = {"Eff NH3-N": "Eff_NH4_N", "Eff NO3-N": "Eff_NO3_N",
           "A1 CH3OH Flow": "A1_FQ_CH3OH", "A2 CH3OH Flow": "A2_FQ_CH3OH",
           "A1 Aeration Flow": "Q_A1_air", "A2 Aeration Flow": "Q_A2_air"}
NAMES = ("eff_nh3_n", "eff_no3_n", "a1_ch3oh_flow", "a2_ch3oh_flow", "a1_air_flow", "a2_air_flow")


def read_workbook(path):
    book = load_workbook(path, data_only=True)
    try:
        sheets = [s for s in book if any(c.value == "Eff NH3-N" for row in s for c in row)]
        if len(sheets) != 1:
            raise ValueError("Expected one sheet with the prediction feature matrix.")
        sheet = sheets[0]
        header = next(row for row in sheet if any(c.value == "Eff NH3-N" for c in row))
        predictions = {str(c.value).strip(): c.column for c in header if c.value in TARGETS}
        if set(predictions) != set(TARGETS):
            raise ValueError("Workbook is missing prediction headings.")
        configs, audit = {}, []
        for name, (heading, target) in zip(NAMES, TARGETS.items()):
            features, marked = [], []
            unit = None
            for row in sheet.iter_rows(min_row=header[0].row + 1):
                alias = row[3].value or row[1].value
                if not row[4].value:
                    continue
                alias = str(alias).strip()
                if alias == target:
                    unit = str(row[2].value or "dimensionless")
                c = row[predictions[heading] - 1]
                checked = str(c.value or "").strip() == "✓"
                green = c.fill.fgColor.type == "rgb" and c.fill.fgColor.rgb == "FFC6EFCE"
                if checked:
                    features.append(alias)
                    if green:
                        marked.append(alias)
                audit.append({"model": name, "heading": heading, "sensor": alias,
                              "checked": checked, "green_target_mark": bool(green),
                              "confirmed_target": target, "cell": c.coordinate})
            if unit is None:
                raise ValueError(f"Target {target} is absent from the sensor rows.")
            configs[name] = {"name": name, "target": target, "features": features,
                             "unit": unit, "heading": heading, "workbook_marked_targets": marked,
                             "target_mark_corrected": marked != [target]}
        return configs, audit
    finally:
        book.close()
